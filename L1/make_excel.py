"""Создаёт excel/variant12.xlsx с моделью для надстройки «Поиск решения» (Solver).

Модель Solver хранится в скрытых именах листа (solver_adj, solver_opt, ...),
поэтому после открытия файла достаточно: Данные -> Поиск решения -> Найти решение.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

from main import A, b, c, signs

SHEET = "Вариант 12"
OUR_X = ["5/2", "0", "9/2", "1/2"]
OUR_Z = "23/2"

bold = Font(bold=True)
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
fill_var = PatternFill("solid", fgColor="FFF2CC")
fill_obj = PatternFill("solid", fgColor="DDEBF7")
center = Alignment(horizontal="center")

wb = Workbook()
ws = wb.active
ws.title = SHEET

ws["A1"] = "Вариант 12: Z = 2x1 + 3x2 + x3 + 4x4 -> min"
ws["A1"].font = Font(bold=True, size=13)

# переменные и цель
ws["A3"] = ""
for k in range(4):
    cell = ws.cell(row=3, column=2 + k, value=f"x{k + 1}")
    cell.font = bold
    cell.alignment = center
ws["A4"] = "Значения"
ws["A5"] = "Коэффициенты Z"
for k in range(4):
    v = ws.cell(row=4, column=2 + k, value=0)
    v.fill = fill_var
    ws.cell(row=5, column=2 + k, value=c[k])
ws["F5"] = "Z ="
ws["F5"].font = bold
ws["G5"] = "=SUMPRODUCT(B5:E5,$B$4:$E$4)"
ws["G5"].fill = fill_obj
ws["G5"].font = bold

# ограничения
ws["A7"] = "Ограничения"
ws["A7"].font = bold
for k, h in enumerate(["x1", "x2", "x3", "x4", "Левая часть", "Знак", "Правая часть"]):
    cell = ws.cell(row=7, column=2 + k, value=h)
    cell.font = bold
    cell.alignment = center
for i, (row, s, v) in enumerate(zip(A, signs, b)):
    r = 8 + i
    ws.cell(row=r, column=1, value=f"Ограничение {i + 1}")
    for k in range(4):
        ws.cell(row=r, column=2 + k, value=row[k])
    ws.cell(row=r, column=6, value=f"=SUMPRODUCT(B{r}:E{r},$B$4:$E$4)")
    ws.cell(row=r, column=7, value=s).alignment = center
    ws.cell(row=r, column=8, value=v)

for r in list(range(3, 6)) + list(range(7, 11)):
    for col in range(1, 9):
        if r in (3, 4, 5) and col > 7:
            continue
        ws.cell(row=r, column=col).border = box

ws["A12"] = "Жёлтые ячейки B4:E4 изменяет Solver, синяя G5 - целевая."
ws["A13"] = "Данные -> Поиск решения -> Найти решение (метод: симплекс-метод, переменные неотрицательны)."

# сравнение с ручным решением
ws["A15"] = "Сравнение с симплекс-таблицами"
ws["A15"].font = bold
for k, h in enumerate(["", "x1", "x2", "x3", "x4", "Z"]):
    cell = ws.cell(row=16, column=1 + k, value=h or None)
    cell.font = bold
    cell.alignment = center
ws["A17"] = "Solver"
ws["A18"] = "Симплекс-метод"
ws["A19"] = "Разница"
for k in range(4):
    col = "BCDE"[k]
    ws[f"{col}17"] = f"={col}4"
    ws[f"{col}18"] = f"={OUR_X[k]}"
    ws[f"{col}19"] = f"=ROUND({col}17-{col}18,6)"
ws["F17"] = "=G5"
ws["F18"] = f"={OUR_Z}"
ws["F19"] = "=ROUND(F17-F18,6)"
ws["A20"] = "Итог"
ws["B20"] = '=IF(AND(F19=0,B19=0,C19=0,D19=0,E19=0),"ответы совпадают","ответы не совпадают")'
ws["B20"].font = bold
for r in range(16, 20):
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = box

ws.column_dimensions["A"].width = 18
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 13

# модель Solver (те же имена, что Excel записывает сам)
ref = f"'{SHEET}'!"
rel = {"<=": 1, "=": 2, ">=": 3}
names = {
    "solver_adj": f"{ref}$B$4:$E$4",
    "solver_opt": f"{ref}$G$5",
    "solver_typ": "2",  # 1 - max, 2 - min, 3 - значение
    "solver_val": "0",
    "solver_eng": "2",  # 2 - симплекс-метод
    "solver_neg": "1",  # переменные без ограничений сделать неотрицательными
    "solver_num": str(len(A)),
    "solver_ver": "3",
    "solver_lin": "1",
    "solver_cvg": "0.0001",
    "solver_drv": "1",
    "solver_est": "1",
    "solver_itr": "2147483647",
    "solver_mip": "2147483647",
    "solver_mni": "30",
    "solver_mrt": "0.075",
    "solver_msl": "2",
    "solver_nod": "2147483647",
    "solver_nwt": "1",
    "solver_pre": "0.000001",
    "solver_rbv": "1",
    "solver_rlx": "2",
    "solver_rsd": "0",
    "solver_scl": "1",
    "solver_sho": "2",
    "solver_ssz": "100",
    "solver_tim": "2147483647",
    "solver_tol": "0.01",
}
for i, s in enumerate(signs, 1):
    names[f"solver_lhs{i}"] = f"{ref}$F${7 + i}"
    names[f"solver_rel{i}"] = str(rel[s])
    names[f"solver_rhs{i}"] = f"{ref}$H${7 + i}"

for name, value in names.items():
    ws.defined_names[name] = DefinedName(name, attr_text=value, hidden=True)

os.makedirs("excel", exist_ok=True)
wb.save("excel/variant12.xlsx")
print("excel/variant12.xlsx")
